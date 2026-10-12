from __future__ import annotations

import json
import hashlib
import re
import sys
import time
from pathlib import Path
from io import BytesIO

import matplotlib.pyplot as plt
import pandas as pd
import streamlit as st
import torch
from PIL import Image

PROJECT = Path(__file__).resolve().parent
if str(PROJECT) not in sys.path:
    sys.path.insert(0, str(PROJECT))

from cnnlab.core import (
    SEED, binary_metrics, build_model, count_parameters, data_root, evaluate_challenge_folder,
    evaluate_internal_test, experiments_dir, list_experiments, load_experiment, make_splits,
    make_transform, processed_dir, project_root, safe_download_microsoft_dataset,
    scan_dataset, train_experiment,
)

st.set_page_config(page_title="CNN Research Lab | CS535", page_icon="🐈", layout="wide")

LANGUAGES = {"Tiếng Việt": "vi", "English": "en", "繁體中文": "zh-TW"}
I18N = {
 "vi": {
  "overview":"Tổng quan", "dataset":"Dataset & QC", "training":"Huấn luyện", "experiments":"Thí nghiệm", "final":"Đánh giá cuối & Challenge", "errors":"Phân tích lỗi & Dự đoán", "survey":"Survey báo cáo", "language":"Ngôn ngữ giao diện", "lab":"Research Lab", "title":"CNN Research Lab", "caption":"CS535 · Thí nghiệm có kiểm soát cho phân loại chó và mèo", "report_title":"Báo cáo survey tổng hợp các thí nghiệm CNN", "report_intro":"Tạo báo cáo Word tổng hợp quy trình, cấu hình, kết quả, learning curves và nhận xét từ các run đã lưu.", "select_runs":"Chọn các lần huấn luyện đưa vào báo cáo", "all_runs":"Chọn tất cả run đã lưu", "generate":"Tạo báo cáo Word (.docx)", "report_note":"Báo cáo chỉ tổng hợp dữ liệu đã lưu. Nếu chưa có run hoặc metric, báo cáo sẽ ghi rõ dữ liệu chưa có; không tự tạo kết quả.", "no_runs":"Chưa có thí nghiệm đã lưu để báo cáo.", "download_ready":"Báo cáo đã được tạo. Tải file Word bên dưới.", "procedure":"Quy trình thực nghiệm", "config":"Cấu hình", "results":"Kết quả", "interpretation":"Nhận xét tự động (dựa trên metric đã lưu)", "limitations":"Giới hạn và lưu ý", "run":"Run", "metric":"Chỉ số", "value":"Giá trị", "language_help":"Áp dụng cho toàn bộ giao diện; tên mô hình, đường dẫn và dữ liệu gốc được giữ nguyên khi cần."},
 "en": {
  "overview":"Overview", "dataset":"Dataset & QC", "training":"Training", "experiments":"Experiments", "final":"Final Evaluation & Challenge", "errors":"Error Analysis & Prediction", "survey":"Survey Report", "language":"Interface language", "lab":"Research Lab", "title":"CNN Research Lab", "caption":"CS535 · Controlled experiments for Cats vs. Dogs classification", "report_title":"CNN Experiments — Consolidated Survey Report", "report_intro":"Generate a Word report summarizing workflow, configurations, results, learning curves, and observations from saved runs.", "select_runs":"Select training runs to include", "all_runs":"Select all saved runs", "generate":"Generate Word report (.docx)", "report_note":"The report only summarizes saved data. Missing runs or metrics are explicitly marked; no results are fabricated.", "no_runs":"No saved experiments are available to report.", "download_ready":"Report generated. Download the Word file below.", "procedure":"Experimental workflow", "config":"Configuration", "results":"Results", "interpretation":"Automated observations (based on saved metrics)", "limitations":"Limitations and notes", "run":"Run", "metric":"Metric", "value":"Value", "language_help":"Applies across the entire interface; model identifiers, paths, and raw data are preserved where needed."},
 "zh-TW": {
  "overview":"總覽", "dataset":"資料集與品質檢查", "training":"模型訓練", "experiments":"實驗比較", "final":"最終評估與挑戰集", "errors":"錯誤分析與預測", "survey":"研究綜合報告", "language":"介面語言", "lab":"研究實驗室", "title":"CNN 研究實驗室", "caption":"CS535 · 貓狗影像分類控制實驗", "report_title":"CNN 實驗綜合研究報告", "report_intro":"將已保存的實驗流程、參數、結果、學習曲線與觀察整理成 Word 報告。", "select_runs":"選擇要納入報告的訓練實驗", "all_runs":"選取所有已保存實驗", "generate":"產生 Word 報告（.docx）", "report_note":"報告僅彙整已保存的資料；缺少的實驗或指標會明確註記，不會虛構結果。", "no_runs":"目前沒有可供報告的已保存實驗。", "download_ready":"報告已產生，請下載下方 Word 檔案。", "procedure":"實驗流程", "config":"參數設定", "results":"實驗結果", "interpretation":"自動觀察（依據已保存指標）", "limitations":"限制與注意事項", "run":"實驗", "metric":"指標", "value":"數值", "language_help":"適用於整個介面；必要時保留模型識別碼、路徑與原始資料。"}
}
def tr(key):
    return I18N.get(st.session_state.get("cnnlab_lang", "vi"), I18N["vi"]).get(key, key)


# Source text -> (Vietnamese, English, Traditional Chinese). This dictionary is
# applied to visible Streamlit text and display-only dataframe/JSON labels.
# Widget values, stored experiment metadata, filenames, and paths are not changed.
_UI_ROWS = [
    ("Chọn phân hệ", "Chọn phân hệ", "Choose a section", "選擇功能區"),
    ("Quản lý dữ liệu, huấn luyện CNN, so sánh thí nghiệm và phân tích lỗi trong cùng một giao diện. Ứng dụng giữ riêng validation và independent test để hỗ trợ quy trình nghiên cứu có thể tái lập.", "Quản lý dữ liệu, huấn luyện CNN, so sánh thí nghiệm và phân tích lỗi trong cùng một giao diện. Ứng dụng giữ riêng validation và independent test để hỗ trợ quy trình nghiên cứu có thể tái lập.", "Manage data, train CNNs, compare experiments, and analyze errors in one interface. Validation and independent test sets are kept separate to support reproducible research.", "在同一介面中管理資料、訓練 CNN、比較實驗並分析錯誤。系統會分開驗證集與獨立測試集，以支援可重現的研究流程。"),
    ("Thiết bị:", "Thiết bị:", "Device:", "運算裝置："),
    ("Seed mặc định:", "Seed mặc định:", "Default seed:", "預設隨機種子："),
    ("Thư mục dự án", "Thư mục dự án", "Project directory", "專案資料夾"),
    ("Bảng điều khiển nghiên cứu", "Bảng điều khiển nghiên cứu", "Research dashboard", "研究儀表板"),
    ("Ảnh hợp lệ trong split", "Ảnh hợp lệ trong split", "Valid images in split", "切分中的有效影像數"),
    ("Chưa chuẩn bị", "Chưa chuẩn bị", "Not prepared", "尚未準備"),
    ("Thí nghiệm đã lưu", "Thí nghiệm đã lưu", "Saved experiments", "已儲存的實驗"),
    ("Run đã đánh giá test", "Run đã đánh giá test", "Runs evaluated on test", "已進行測試集評估的實驗"),
    ("Thiết bị", "Thiết bị", "Device", "裝置"),
    ("Quy trình thực nghiệm", "Quy trình thực nghiệm", "Experimental workflow", "實驗流程"),
    ("**Dataset → QC/manifest → stratified split 70/15/15 → baseline/controlled runs → chọn cấu hình bằng validation → independent test một lần sau khi khóa mô hình → error analysis.**", "**Dataset → QC/manifest → stratified split 70/15/15 → chọn cấu hình bằng validation → khóa mô hình → independent test → phân tích lỗi.**", "**Dataset → QC/manifest → stratified 70/15/15 split → baseline/controlled runs → select configuration using validation → lock the model → one independent test → error analysis.**", "**資料集 → 品質檢查/清單 → 分層切分 70/15/15 → 基準與控制實驗 → 依驗證集選擇設定 → 鎖定模型 → 執行一次獨立測試 → 錯誤分析。**"),
    ("Không chọn mô hình hoặc chỉnh siêu tham số dựa trên independent test. Trong giai đoạn phát triển, hãy dùng validation metrics và validation errors.", "Không chọn mô hình hoặc chỉnh siêu tham số dựa trên independent test. Trong giai đoạn phát triển, hãy dùng validation metrics và validation errors.", "Do not select models or tune hyperparameters using the independent test set. During development, use validation metrics and validation errors.", "不得依據獨立測試集選擇模型或調整超參數。開發階段請使用驗證指標與驗證錯誤。"),
    ("Các bước bắt đầu", "Các bước bắt đầu", "Getting started", "開始使用"),
    ("1. Mở **Dataset & QC** để tải dataset Microsoft hoặc quét thư mục ảnh sẵn có.\n2. Tạo manifest và split tái lập.\n3. Mở **Huấn luyện** để chạy baseline CNN hoặc kiến trúc TorchVision.\n4. So sánh các run tại **Thí nghiệm**; khi đã khóa lựa chọn, mới mở **Đánh giá cuối & Challenge**.", "1. Mở **Dataset & QC** để tải dataset Microsoft hoặc quét thư mục ảnh sẵn có.\n2. Tạo manifest và split tái lập.\n3. Mở **Huấn luyện** để chạy baseline CNN hoặc kiến trúc TorchVision.\n4. So sánh các run tại **Thí nghiệm**; khi đã khóa lựa chọn, mới mở **Đánh giá cuối & Challenge**.", "1. Open **Dataset & QC** to download the Microsoft dataset or scan an existing image folder.\n2. Create a reproducible manifest and split.\n3. Open **Training** to run a baseline CNN or a TorchVision architecture.\n4. Compare runs under **Experiments**; open **Final Evaluation & Challenge** only after locking the choice.", "1. 前往 **資料集與品質檢查** 下載 Microsoft 資料集或掃描現有影像資料夾。\n2. 建立可重現的清單與資料切分。\n3. 前往 **模型訓練** 執行基準 CNN 或 TorchVision 架構。\n4. 在 **實驗比較** 中比較各次實驗；鎖定模型後才進行 **最終評估與挑戰集**。"),
    ("Các run gần đây", "Các run gần đây", "Recent runs", "最近的實驗"),
    ("Chưa có run nào. Hãy chuẩn bị dataset trước khi huấn luyện.", "Chưa có run nào. Hãy chuẩn bị dataset trước khi huấn luyện.", "No runs yet. Prepare a dataset before training.", "目前沒有實驗。請先準備資料集再開始訓練。"),
    ("Quản lý dataset và kiểm tra chất lượng", "Quản lý dataset và kiểm tra chất lượng", "Dataset management and quality checks", "資料集管理與品質檢查"),
    ("Nguồn chính được khuyến nghị trong tài liệu CS535 là Microsoft/Asirra Cats vs. Dogs (~25.000 ảnh thô).", "Nguồn chính được khuyến nghị trong tài liệu CS535 là Microsoft/Asirra Cats vs. Dogs (~25.000 ảnh thô).", "The primary dataset recommended in the CS535 materials is Microsoft/Asirra Cats vs. Dogs (about 25,000 raw images).", "CS535 教材建議的主要資料集為 Microsoft/Asirra Cats vs. Dogs（約 25,000 張原始影像）。"),
    ("Nguồn chính thức: [Microsoft Cats and Dogs](https://www.microsoft.com/en-us/download/details.aspx?id=54765)", "Nguồn chính thức: [Microsoft Cats and Dogs](https://www.microsoft.com/en-us/download/details.aspx?id=54765)", "Official source: [Microsoft Cats and Dogs](https://www.microsoft.com/en-us/download/details.aspx?id=54765)", "官方來源：[Microsoft Cats and Dogs](https://www.microsoft.com/en-us/download/details.aspx?id=54765)"),
    ("Thư mục dataset", "Thư mục dataset", "Dataset directory", "資料集資料夾"),
    ("Có thể là thư mục chứa PetImages, hoặc chính thư mục có hai thư mục Cat và Dog.", "Có thể là thư mục chứa PetImages, hoặc chính thư mục có hai thư mục Cat và Dog.", "Use the parent directory containing PetImages, or the directory containing the Cat and Dog subfolders.", "可填入包含 PetImages 的上層資料夾，或直接填入包含 Cat 與 Dog 子資料夾的目錄。"),
    ("Tải dataset Microsoft (~786 MB)", "Tải dataset Microsoft (~786 MB)", "Download Microsoft dataset (~786 MB)", "下載 Microsoft 資料集（約 786 MB）"),
    ("Quét ảnh, loại file lỗi/trùng và tạo split", "Quét ảnh, loại file lỗi/trùng và tạo split", "Scan images, remove corrupt/duplicate files, and create splits", "掃描影像、移除損壞或重複檔案並建立資料切分"),
    ("Đang quét ảnh và kiểm tra file. Với khoảng 25.000 ảnh, bước này có thể mất vài phút.", "Đang quét ảnh và kiểm tra file. Với khoảng 25.000 ảnh, bước này có thể mất vài phút.", "Scanning images and checking files. With about 25,000 images, this may take a few minutes.", "正在掃描影像並檢查檔案。約 25,000 張影像可能需要幾分鐘。"),
    ("Đã tải/giải nén dữ liệu trong", "Đã tải/giải nén dữ liệu trong", "Dataset downloaded/extracted to", "資料集已下載／解壓縮至"),
    ("Không thể tải hoặc giải nén:", "Không thể tải hoặc giải nén:", "Download or extraction failed:", "下載或解壓縮失敗："),
    ("Quét và chia tập dữ liệu hoàn tất.", "Quét và chia tập dữ liệu hoàn tất.", "Dataset scan and split completed.", "資料集掃描與切分已完成。"),
    ("Không thể chuẩn bị dataset:", "Không thể chuẩn bị dataset:", "Could not prepare dataset:", "無法準備資料集："),
    ("Báo cáo kiểm tra dữ liệu gần nhất", "Báo cáo kiểm tra dữ liệu gần nhất", "Latest dataset QC report", "最近一次資料集品質檢查報告"),
    ("Ảnh hợp lệ sau dedup", "Ảnh hợp lệ sau dedup", "Valid images after deduplication", "去除重複後的有效影像"),
    ("Cat", "Mèo", "Cat", "貓"),
    ("Dog", "Chó", "Dog", "狗"),
    ("File lỗi / bỏ qua", "File lỗi / bỏ qua", "Invalid / skipped files", "無效／略過的檔案"),
    ("Thống kê kích thước ảnh nguồn", "Thống kê kích thước ảnh nguồn", "Source image dimension statistics", "來源影像尺寸統計"),
    ("Exact duplicate đã loại:", "Exact duplicate đã loại:", "Exact duplicates removed:", "已移除的完全重複影像："),
    ("Trùng ảnh nhưng mâu thuẫn nhãn đã loại:", "Trùng ảnh nhưng mâu thuẫn nhãn đã loại:", "Cross-class label conflicts removed:", "已移除的跨類別標籤衝突影像："),
    ("Quét lúc:", "Quét lúc:", "Scanned at:", "掃描時間："),
    ("Một số file không hợp lệ (tối đa 100 mục được lưu)", "Một số file không hợp lệ (tối đa 100 mục được lưu)", "Some invalid files (up to 100 saved entries)", "部分無效檔案（最多保存 100 筆）"),
    ("Phân chia train / validation / independent test", "Phân chia train / validation / independent test", "Train / validation / independent test split", "訓練／驗證／獨立測試集切分"),
    ("Seed = 42; stratified 70/15/15. Split đã được lưu để các thí nghiệm dùng cùng một phân chia.", "Seed = 42; stratified 70/15/15. Split đã được lưu để các thí nghiệm dùng cùng một phân chia.", "Seed = 42; stratified 70/15/15. The split is saved so all experiments use the same partition.", "隨機種子為 42；分層切分比例為 70/15/15。切分結果已保存，所有實驗會使用相同的資料分組。"),
    ("Chưa có split manifest. Quét dataset để hệ thống tạo manifest và chia tập.", "Chưa có split manifest. Quét dataset để hệ thống tạo manifest và chia tập.", "No split manifest found. Scan the dataset to create a manifest and split.", "找不到資料切分清單。請掃描資料集以建立清單與資料切分。"),
    ("Dataset đánh giá ngoài: Oxford-IIIT Pet", "Dataset đánh giá ngoài: Oxford-IIIT Pet", "External evaluation dataset: Oxford-IIIT Pet", "外部評估資料集：Oxford-IIIT Pet"),
    ("Theo đề xuất nghiên cứu, có thể giữ Oxford-IIIT Pet làm external test độc lập. MVP này không tự tải/chuyển nhãn Oxford; hãy chuẩn bị thư mục chuẩn hóa cat/ và dog/ để đánh giá ở trang Challenge.", "Theo đề xuất nghiên cứu, có thể giữ Oxford-IIIT Pet làm external test độc lập. MVP này không tự tải/chuyển nhãn Oxford; hãy chuẩn bị thư mục chuẩn hóa cat/ và dog/ để đánh giá ở trang Challenge.", "Oxford-IIIT Pet can be kept as an independent external test set. This MVP does not download or remap Oxford labels automatically; prepare normalized cat/ and dog/ folders for evaluation on the Challenge page.", "可將 Oxford-IIIT Pet 保留為獨立外部測試集。此 MVP 不會自動下載或轉換 Oxford 標籤；請準備標準化的 cat/ 與 dog/ 資料夾，再至挑戰集頁面評估。"),
    ("[Oxford-IIIT Pet dataset](https://www.robots.ox.ac.uk/~vgg/data/pets/)", "[Oxford-IIIT Pet dataset](https://www.robots.ox.ac.uk/~vgg/data/pets/)", "[Oxford-IIIT Pet dataset](https://www.robots.ox.ac.uk/~vgg/data/pets/)", "[Oxford-IIIT Pet 資料集](https://www.robots.ox.ac.uk/~vgg/data/pets/)"),
    ("Cấu hình một thí nghiệm", "Cấu hình một thí nghiệm", "Configure an experiment", "設定實驗"),
    ("Chưa có dataset split. Vào Dataset & QC để quét dữ liệu trước.", "Chưa có dataset split. Vào Dataset & QC để quét dữ liệu trước.", "No dataset split found. Go to Dataset & QC and scan the data first.", "尚未建立資料切分。請先前往資料集與品質檢查頁面掃描資料。"),
    ("Kiến trúc", "Kiến trúc", "Architecture", "模型架構"),
    ("Image resolution", "Kích thước ảnh đầu vào", "Image resolution", "輸入影像尺寸"),
    ("Batch size", "Kích thước batch", "Batch size", "批次大小"),
    ("Epochs", "Số epoch", "Epochs", "訓練 Epoch 數"),
    ("Optimizer", "Bộ tối ưu", "Optimizer", "最佳化器"),
    ("Learning rate", "Tốc độ học", "Learning rate", "學習率"),
    ("Weight decay", "Hệ số weight decay", "Weight decay", "權重衰減"),
    ("Train images per run", "Số ảnh train mỗi run", "Training images per run", "每次實驗使用的訓練影像數"),
    ("Toàn bộ train split", "Toàn bộ train split", "Entire training split", "完整訓練集"),
    ("Tối đa", "Tối đa", "Up to", "最多"),
    ("quick run", "chạy nhanh", "quick run", "快速實驗"),
    ("Data augmentation nhẹ", "Tăng cường dữ liệu mức nhẹ", "Light data augmentation", "輕量資料增強"),
    ("Transfer-learning strategy", "Chiến lược transfer learning", "Transfer-learning strategy", "遷移學習策略"),
    ("From scratch", "Huấn luyện từ đầu", "From scratch", "從頭訓練"),
    ("Feature extraction (freeze backbone)", "Trích xuất đặc trưng (đóng băng backbone)", "Feature extraction (freeze backbone)", "特徵擷取（凍結主幹網路）"),
    ("Fine-tuning (pretrained backbone)", "Fine-tuning (tinh chỉnh backbone pretrained)", "Fine-tuning (pretrained backbone)", "微調（使用預訓練主幹網路）"),
    ("Dùng với kiến trúc TorchVision. Feature extraction chỉ cập nhật classifier; fine-tuning cập nhật toàn bộ model.", "Dùng với kiến trúc TorchVision. Feature extraction chỉ cập nhật classifier; fine-tuning cập nhật toàn bộ model.", "Used with TorchVision architectures. Feature extraction updates only the classifier; fine-tuning updates the entire model.", "適用於 TorchVision 架構。特徵擷取只更新分類器；微調會更新整個模型。"),
    ("Fine-tuning thường nên bắt đầu với learning rate nhỏ (ví dụ 1e-4) để tránh phá vỡ trọng số pretrained.", "Fine-tuning thường nên bắt đầu với learning rate nhỏ (ví dụ 1e-4) để tránh phá vỡ trọng số pretrained.", "Fine-tuning usually starts with a small learning rate (e.g., 1e-4) to avoid disrupting pretrained weights.", "微調通常應從較小的學習率開始（例如 1e-4），避免破壞預訓練權重。"),
    ("Pretrained weights có thể cần tải từ Internet ở lần chạy đầu. Mọi kiến trúc đều dùng cùng split đã lưu.", "Pretrained weights có thể cần tải từ Internet ở lần chạy đầu. Mọi kiến trúc đều dùng cùng split đã lưu.", "Pretrained weights may need to be downloaded on the first run. All architectures use the same saved split.", "首次執行時可能需要從網路下載預訓練權重。所有架構都使用相同的已保存切分。"),
    ("**Tham số riêng của Configurable CNN**", "**Tham số riêng của Configurable CNN**", "**Configurable CNN-specific parameters**", "**可設定 CNN 的專屬參數**"),
    ("Depth / convolution blocks", "Độ sâu / số khối tích chập", "Depth / convolution blocks", "深度／卷積區塊數"),
    ("Base filters", "Số filter cơ sở", "Base filters", "基礎濾波器數量"),
    ("Kernel size", "Kích thước kernel", "Kernel size", "核心大小"),
    ("Pooling", "Phương pháp pooling", "Pooling", "池化方式"),
    ("max", "max", "max pooling", "最大池化"),
    ("avg", "avg", "average pooling", "平均池化"),
    ("none", "không dùng", "none", "不使用"),
    ("Dropout", "Dropout", "Dropout", "Dropout"),
    ("Batch normalization", "Chuẩn hóa batch", "Batch normalization", "批次正規化"),
    ("Activation", "Hàm kích hoạt", "Activation", "啟動函數"),
    ("Random seed", "Hạt giống ngẫu nhiên", "Random seed", "隨機種子"),
    ("Bắt đầu huấn luyện", "Bắt đầu huấn luyện", "Start training", "開始訓練"),
    ("Đang khởi tạo model và huấn luyện...", "Đang khởi tạo model và huấn luyện...", "Initializing model and training...", "正在初始化模型並開始訓練……"),
    ("Epoch", "Epoch", "Epoch", "訓練輪次"),
    ("loss train=", "loss huấn luyện=", "training loss=", "訓練損失="),
    ("loss val=", "loss kiểm định=", "validation loss=", "驗證損失="),
    ("Validation accuracy:", "Độ chính xác validation:", "Validation accuracy:", "驗證準確率："),
    ("Validation ROC-AUC:", "ROC-AUC validation:", "Validation ROC-AUC:", "驗證 ROC-AUC："),
    ("thời gian epoch:", "thời gian mỗi epoch:", "epoch time:", "每輪訓練時間："),
    ("Đã lưu thí nghiệm:", "Đã lưu thí nghiệm:", "Experiment saved:", "實驗已儲存："),
    ("Checkpoint, config, history, validation metrics và validation predictions đã được lưu trong thư mục run.", "Checkpoint, config, history, validation metrics và validation predictions đã được lưu trong thư mục run.", "Checkpoint, configuration, training history, validation metrics, and validation predictions were saved in the run directory.", "模型檢查點、設定、訓練歷程、驗證指標與驗證預測已儲存於實驗資料夾。"),
    ("Validation accuracy", "Độ chính xác validation", "Validation accuracy", "驗證準確率"),
    ("Validation ROC-AUC", "ROC-AUC validation", "Validation ROC-AUC", "驗證 ROC-AUC"),
    ("Validation F1", "F1 validation", "Validation F1", "驗證 F1"),
    ("Trainable parameters", "Số tham số có thể huấn luyện", "Trainable parameters", "可訓練參數數量"),
    ("Các chỉ số đang hiển thị là validation metrics; internal test chưa được sử dụng.", "Các chỉ số đang hiển thị là validation metrics; internal test chưa được sử dụng.", "The displayed values are validation metrics; the internal test set has not been used.", "目前顯示的是驗證指標；尚未使用內部測試集。"),
    ("Huấn luyện thất bại:", "Huấn luyện thất bại:", "Training failed:", "訓練失敗："),
    ("Experiment tracking & comparison", "Theo dõi và so sánh thí nghiệm", "Experiment tracking & comparison", "實驗追蹤與比較"),
    ("Chưa có thí nghiệm. Hãy chạy ít nhất một baseline.", "Chưa có thí nghiệm. Hãy chạy ít nhất một baseline.", "No experiments yet. Run at least one baseline.", "目前沒有實驗。請至少執行一次基準實驗。"),
    ("Tải bảng so sánh CSV", "Tải bảng so sánh CSV", "Download comparison table (CSV)", "下載比較表（CSV）"),
    ("So sánh validation metrics", "So sánh các chỉ số validation", "Compare validation metrics", "比較驗證指標"),
    ("Chỉ số", "Chỉ số", "Metric", "指標"),
    ("Learning curve", "Đường cong học tập", "Learning curve", "學習曲線"),
    ("Chọn run để xem learning curve", "Chọn run để xem learning curve", "Select a run to view learning curves", "選擇要檢視學習曲線的實驗"),
    ("Train loss", "Loss train", "Training loss", "訓練損失"),
    ("Validation loss", "Loss validation", "Validation loss", "驗證損失"),
    ("Epoch", "Epoch", "Epoch", "訓練輪次"),
    ("BCE loss", "BCE loss", "BCE loss", "BCE 損失"),
    ("Sử dụng validation để so sánh. Không xếp hạng/chọn run theo internal test metrics.", "Sử dụng validation để so sánh. Không xếp hạng/chọn run theo internal test metrics.", "Use validation for comparisons. Do not rank or select runs using internal test metrics.", "請使用驗證集進行比較，不要依據內部測試指標排名或選擇實驗。"),
    ("Chi tiết cấu hình của run đã chọn", "Chi tiết cấu hình của run đã chọn", "Selected run configuration details", "所選實驗的設定詳細資料"),
    ("Independent test và external challenge", "Independent test và external challenge", "Independent test and external challenge", "獨立測試與外部挑戰集"),
    ("Chỉ chạy internal test sau khi đã khóa mô hình và siêu tham số. Không dùng kết quả này để quay lại chọn run; mỗi lần đánh giá test đều nên được ghi nhận trong báo cáo nghiên cứu.", "Chỉ chạy internal test sau khi đã khóa mô hình và siêu tham số. Không dùng kết quả này để quay lại chọn run; mỗi lần đánh giá test đều nên được ghi nhận trong báo cáo nghiên cứu.", "Run the internal test only after locking the model and hyperparameters. Do not use these results to choose another run; record every test evaluation in the research report.", "只有在鎖定模型與超參數後才能執行內部測試。不得利用測試結果回頭選擇其他實驗；每次測試評估都應記錄在研究報告中。"),
    ("Chưa có checkpoint để đánh giá.", "Chưa có checkpoint để đánh giá.", "No checkpoint is available for evaluation.", "沒有可供評估的模型檢查點。"),
    ("Chọn checkpoint đã chốt để đánh giá", "Chọn checkpoint đã chốt để đánh giá", "Select the locked checkpoint for evaluation", "選擇已鎖定的模型檢查點進行評估"),
    ("Run này không được huấn luyện trên split manifest hiện tại. Không thể dùng internal test hiện tại để đánh giá run này một cách hợp lệ.", "Run này không được huấn luyện trên split manifest hiện tại. Không thể dùng internal test hiện tại để đánh giá run này một cách hợp lệ.", "This run was not trained on the current split manifest. The current internal test cannot validly evaluate this run.", "此實驗並非使用目前的資料切分清單訓練，無法以目前的內部測試集進行有效評估。"),
    ("Independent test đã được sử dụng cho run", "Independent test đã được sử dụng cho run", "The independent test has already been used for run", "獨立測試集已用於實驗"),
    ("Theo quy trình nghiên cứu, không chạy lại test trên các run khác để chọn mô hình.", "Theo quy trình nghiên cứu, không chạy lại test trên các run khác để chọn mô hình.", "Under the research protocol, do not rerun the test on other runs to select a model.", "依照研究流程，不應在其他實驗上重複測試以選擇模型。"),
    ("Run này có kết quả internal test đã lưu. Kết quả này chỉ dành cho báo cáo cuối.", "Run này có kết quả internal test đã lưu. Kết quả này chỉ dành cho báo cáo cuối.", "This run has saved internal-test results. These results are for final reporting only.", "此實驗已有保存的內部測試結果，僅供最終報告使用。"),
    ("Independent test đã khóa", "Independent test đã khóa", "Independent test locked", "獨立測試已鎖定"),
    ("Chạy internal test cho checkpoint đã chốt", "Chạy internal test cho checkpoint đã chốt", "Run internal test on the locked checkpoint", "對已鎖定的模型執行內部測試"),
    ("Không tìm thấy split manifest.", "Không tìm thấy split manifest.", "Split manifest not found.", "找不到資料切分清單。"),
    ("Đang chạy independent test...", "Đang chạy independent test...", "Running independent test...", "正在執行獨立測試……"),
    ("Đã lưu đánh giá independent test với n=", "Đã lưu đánh giá independent test với n=", "Independent-test evaluation saved with n=", "獨立測試評估已儲存，樣本數 n="),
    ("Test evaluation thất bại:", "Test evaluation thất bại:", "Test evaluation failed:", "測試評估失敗："),
    ("External / hard challenge set", "Bộ challenge ngoài / khó", "External / hard challenge set", "外部／高難度挑戰集"),
    ("Chuẩn bị thư mục ngoài dataset theo cấu trúc: `challenge_set/cat/`, `challenge_set/dog/`, `challenge_set/ood/`. Nhãn cat/dog dùng để tính metric; OOD chỉ phân tích confidence, không tính binary accuracy.", "Chuẩn bị thư mục ngoài dataset theo cấu trúc: `challenge_set/cat/`, `challenge_set/dog/`, `challenge_set/ood/`. Nhãn cat/dog dùng để tính metric; OOD chỉ phân tích confidence, không tính binary accuracy.", "Prepare an external dataset with `challenge_set/cat/`, `challenge_set/dog/`, and `challenge_set/ood/`. Cat/dog labels are used for metrics; OOD is analyzed for confidence only, not binary accuracy.", "請依照 `challenge_set/cat/`、`challenge_set/dog/` 與 `challenge_set/ood/` 的結構準備外部資料集。cat/dog 標籤用於計算指標；OOD 僅分析信心分數，不計算二元準確率。"),
    ("Đường dẫn challenge_set", "Đường dẫn challenge_set", "Challenge-set path", "挑戰集路徑"),
    ("Chỉ đánh giá external challenge sau khi đã khóa mô hình cuối qua internal test. Dùng validation để chọn mô hình trước bước đó.", "Chỉ đánh giá external challenge sau khi đã khóa mô hình cuối qua internal test. Dùng validation để chọn mô hình trước bước đó.", "Evaluate the external challenge set only after locking the final model through internal testing. Use validation to choose the model beforehand.", "只有在內部測試後鎖定最終模型，才能評估外部挑戰集；在此之前請使用驗證集選擇模型。"),
    ("Đánh giá challenge folder", "Đánh giá challenge folder", "Evaluate challenge folder", "評估挑戰資料夾"),
    ("Đã đánh giá", "Đã đánh giá", "Evaluated", "已評估"),
    ("ảnh challenge; không cập nhật trọng số mô hình.", "ảnh challenge; không cập nhật trọng số mô hình.", "challenge images; model weights were not updated.", "張挑戰影像；未更新模型權重。"),
    ("Challenge evaluation thất bại:", "Challenge evaluation thất bại:", "Challenge evaluation failed:", "挑戰集評估失敗："),
    ("Kết quả challenge đã lưu", "Kết quả challenge đã lưu", "Saved challenge results", "已保存的挑戰集結果"),
    ("Phân tích lỗi và dự đoán ảnh mới", "Phân tích lỗi và dự đoán ảnh mới", "Error analysis and new-image prediction", "錯誤分析與新影像預測"),
    ("Chưa có checkpoint. Huấn luyện một run trước.", "Chưa có checkpoint. Huấn luyện một run trước.", "No checkpoint found. Train a run first.", "找不到模型檢查點。請先訓練一次。"),
    ("Checkpoint", "Checkpoint", "Checkpoint", "模型檢查點"),
    ("Validation (khuyến nghị trong phát triển)", "Validation (khuyến nghị trong phát triển)", "Validation (recommended during development)", "驗證集（開發階段建議使用）"),
    ("Internal test (chỉ xem báo cáo cuối) ", "Internal test (chỉ xem báo cáo cuối) ", "Internal test (final report only)", "內部測試集（僅供最終報告）"),
    ("External / hard challenge", "External / hard challenge", "External / hard challenge", "外部／高難度挑戰集"),
    ("Nguồn phân tích lỗi", "Nguồn phân tích lỗi", "Error-analysis source", "錯誤分析來源"),
    ("Số ảnh trong file", "Số ảnh trong file", "Images in file", "檔案中的影像數"),
    ("Lỗi dự đoán", "Lỗi dự đoán", "Prediction errors", "預測錯誤數"),
    ("Tỷ lệ lỗi", "Tỷ lệ lỗi", "Error rate", "錯誤率"),
    ("**Các trường hợp dự đoán sai có confidence cao**", "**Các trường hợp dự đoán sai có confidence cao**", "**High-confidence misclassifications**", "**高信心度的錯誤分類案例**"),
    ("Số ảnh lỗi để xem", "Số ảnh lỗi để xem", "Number of error images to display", "要顯示的錯誤影像數"),
    ("Không mở được ảnh:", "Không mở được ảnh:", "Could not open image:", "無法開啟影像："),
    ("Không tìm thấy trường hợp dự đoán sai trong file đã chọn, hoặc nguồn dữ liệu không có nhãn.", "Không tìm thấy trường hợp dự đoán sai trong file đã chọn, hoặc nguồn dữ liệu không có nhãn.", "No misclassified examples were found in the selected file, or the data source has no labels.", "所選檔案中沒有錯誤分類案例，或資料來源沒有標籤。"),
    ("Tải prediction table CSV", "Tải prediction table CSV", "Download prediction table (CSV)", "下載預測表（CSV）"),
    ("Chưa có prediction file cho nguồn này.", "Chưa có prediction file cho nguồn này.", "No prediction file is available for this source.", "此來源沒有可用的預測檔案。"),
    ("Thử dự đoán một ảnh mới", "Thử dự đoán một ảnh mới", "Try predicting a new image", "試著預測新影像"),
    ("Tải ảnh JPG / PNG / WEBP", "Tải ảnh JPG / PNG / WEBP", "Upload an image (JPG / PNG / WEBP)", "上傳影像（JPG／PNG／WEBP）"),
    ("Ảnh đầu vào", "Ảnh đầu vào", "Input image", "輸入影像"),
    ("Phân loại ảnh", "Phân loại ảnh", "Classify image", "分類影像"),
    ("Dự đoán", "Dự đoán", "Prediction", "預測結果"),
    ("Xác suất là đầu ra sigmoid của bộ phân loại nhị phân; không tự động đồng nghĩa với xác suất đã hiệu chuẩn.", "Xác suất là đầu ra sigmoid của bộ phân loại nhị phân; không tự động đồng nghĩa với xác suất đã hiệu chuẩn.", "The probability is the sigmoid output of a binary classifier; it is not necessarily a calibrated probability.", "此機率為二元分類器的 sigmoid 輸出，不一定代表已校準的機率。"),
    ("Không thể dự đoán:", "Không thể dự đoán:", "Prediction failed:", "無法進行預測："),
    ("Run ID", "Mã run", "Run ID", "實驗 ID"),
    ("run_id", "Mã run", "Run ID", "實驗 ID"),
    ("model_name", "Tên mô hình", "Model name", "模型名稱"),
    ("created_at", "Thời điểm tạo", "Created at", "建立時間"),
    ("train_images_used", "Số ảnh train đã dùng", "Training images used", "使用的訓練影像數"),
    ("val_accuracy", "Độ chính xác validation", "Validation accuracy", "驗證準確率"),
    ("val_balanced_accuracy", "Độ chính xác cân bằng validation", "Validation balanced accuracy", "驗證平衡準確率"),
    ("val_precision", "Precision validation", "Validation precision", "驗證精確率"),
    ("val_recall_sensitivity", "Recall / sensitivity validation", "Validation recall / sensitivity", "驗證召回率／敏感度"),
    ("val_specificity", "Specificity validation", "Validation specificity", "驗證特異度"),
    ("val_f1", "F1 validation", "Validation F1", "驗證 F1"),
    ("val_roc_auc", "ROC-AUC validation", "Validation ROC-AUC", "驗證 ROC-AUC"),
    ("val_pr_auc_avg_precision", "PR-AUC / average precision validation", "Validation PR-AUC / average precision", "驗證 PR-AUC／平均精確率"),
    ("elapsed_seconds", "Thời gian chạy (giây)", "Elapsed time (seconds)", "執行時間（秒）"),
    ("parameters", "Số tham số", "Parameter count", "參數數量"),
    ("image_size", "Kích thước ảnh", "Image size", "影像尺寸"),
    ("batch_size", "Kích thước batch", "Batch size", "批次大小"),
    ("epochs", "Số epoch", "Epochs", "Epoch 數"),
    ("optimizer", "Bộ tối ưu", "Optimizer", "最佳化器"),
    ("learning_rate", "Tốc độ học", "Learning rate", "學習率"),
    ("weight_decay", "Hệ số weight decay", "Weight decay", "權重衰減"),
    ("augmentation", "Tăng cường dữ liệu", "Data augmentation", "資料增強"),
    ("pretrained", "Dùng trọng số pretrained", "Pretrained weights", "使用預訓練權重"),
    ("transfer_strategy", "Chiến lược transfer learning", "Transfer-learning strategy", "遷移學習策略"),
    ("depth", "Độ sâu CNN", "CNN depth", "CNN 深度"),
    ("base_filters", "Số filter cơ sở", "Base filters", "基礎濾波器數量"),
    ("kernel_size", "Kích thước kernel", "Kernel size", "核心大小"),
    ("pool", "Pooling", "Pooling", "池化"),
    ("dropout", "Dropout", "Dropout", "Dropout"),
    ("batchnorm", "Batch normalization", "Batch normalization", "批次正規化"),
    ("activation", "Hàm kích hoạt", "Activation", "啟動函數"),
    ("split_ratio", "Tỷ lệ chia tập", "Split ratio", "資料切分比例"),
    ("accuracy", "Độ chính xác", "Accuracy", "準確率"),
    ("balanced_accuracy", "Độ chính xác cân bằng", "Balanced accuracy", "平衡準確率"),
    ("precision", "Precision", "Precision", "精確率"),
    ("recall_sensitivity", "Recall / sensitivity", "Recall / sensitivity", "召回率／敏感度"),
    ("specificity", "Specificity", "Specificity", "特異度"),
    ("f1", "F1", "F1", "F1"),
    ("roc_auc", "ROC-AUC", "ROC-AUC", "ROC-AUC"),
    ("pr_auc_avg_precision", "PR-AUC / average precision", "PR-AUC / average precision", "PR-AUC／平均精確率"),
    ("mcc", "MCC", "MCC", "MCC"),
    ("brier", "Brier score", "Brier score", "Brier 分數"),
    ("n_samples", "Số mẫu", "Number of samples", "樣本數"),
    ("n_images", "Số ảnh", "Number of images", "影像數"),
    ("path", "Đường dẫn", "Path", "路徑"),
    ("label", "Nhãn", "Label", "標籤"),
    ("split", "Tập dữ liệu", "Split", "資料分組"),
    ("class_name", "Tên lớp", "Class name", "類別名稱"),
    ("width", "Chiều rộng", "Width", "寬度"),
    ("height", "Chiều cao", "Height", "高度"),
    ("min_side", "Cạnh ngắn", "Minimum side", "短邊"),
    ("max_side", "Cạnh dài", "Maximum side", "長邊"),
    ("aspect_ratio", "Tỷ lệ khung hình", "Aspect ratio", "長寬比"),
    ("Learning history", "Lịch sử huấn luyện", "Learning history", "訓練歷程"),
    ("Train loss", "Loss train", "Training loss", "訓練損失"),
    ("Validation loss", "Loss validation", "Validation loss", "驗證損失"),
    ("P(Dog)", "P(Chó)", "P(Dog)", "P(狗)"),
    ("P(Cat)", "P(Mèo)", "P(Cat)", "P(貓)"),
    ("Không thể tạo báo cáo Word:", "Không thể tạo báo cáo Word:", "Could not create Word report:", "無法建立 Word 報告："),
    ("Model name", "Tên mô hình", "Model name", "模型名稱"),
    ("Created", "Thời điểm tạo", "Created", "建立時間"),
    ("Val Accuracy", "Độ chính xác validation", "Validation accuracy", "驗證準確率"),
    ("Val F1", "F1 validation", "Validation F1", "驗證 F1"),
    ("Val ROC-AUC", "ROC-AUC validation", "Validation ROC-AUC", "驗證 ROC-AUC"),
    ("epoch", "epoch", "epoch", "訓練輪次"),
    ("train_loss", "Loss train", "Training loss", "訓練損失"),
    ("val_loss", "Loss validation", "Validation loss", "驗證損失"),
    ("Validation F1 =", "F1 validation =", "Validation F1 =", "驗證 F1 ="),
    ("Validation ROC-AUC =", "ROC-AUC validation =", "Validation ROC-AUC =", "驗證 ROC-AUC ="),
    ("Validation accuracy =", "Độ chính xác validation =", "Validation accuracy =", "驗證準確率 ="),
    ("Training loss changed from", "Loss train thay đổi từ", "Training loss changed from", "訓練損失從"),
    ("Final validation loss is notably higher than training loss; inspect possible overfitting.", "Loss validation cuối cao hơn đáng kể so với loss train; cần kiểm tra khả năng overfitting.", "Final validation loss is notably higher than training loss; inspect possible overfitting.", "最後的驗證損失明顯高於訓練損失，請檢查是否過度擬合。"),
    ("Compare the training and validation curves to assess generalization.", "So sánh đường cong train và validation để đánh giá khả năng tổng quát hóa.", "Compare the training and validation curves to assess generalization.", "請比較訓練與驗證曲線，以評估泛化能力。"),
    ("Insufficient saved metrics for an evidence-based observation.", "Chưa đủ metric được lưu để đưa ra nhận xét có căn cứ.", "Insufficient saved metrics for an evidence-based observation.", "目前保存的指標不足，無法提出有依據的觀察。"),
    ("Generated:", "Thời điểm tạo:", "Generated:", "產生時間："),
    ("Important:", "Lưu ý:", "Important:", "重要："),
    ("True:", "Nhãn thật:", "True label:", "真實標籤："),
    ("Pred:", "Dự đoán:", "Predicted:", "預測："),
    ("p(Dog)=", "p(Chó)=", "p(Dog)=", "p(狗)="),
    ("val AUC", "AUC validation", "validation AUC", "驗證 AUC"),
    ("Configurable CNN (from scratch)", "CNN tùy chỉnh (huấn luyện từ đầu)", "Configurable CNN (from scratch)", "可設定 CNN（從頭訓練）"),
    ("CUDA GPU", "GPU CUDA", "CUDA GPU", "CUDA GPU"),
    ("GPU", "GPU", "GPU", "GPU"),
    ("CPU", "CPU", "CPU", "CPU"),
]

_UI_TRANSLATIONS = {
    source: {"vi": vi, "en": en, "zh-TW": zh}
    for source, vi, en, zh in _UI_ROWS
}


def translate_ui(value):
    """Translate display text without changing underlying widget/data values."""
    if not isinstance(value, str):
        return value
    lang = st.session_state.get("cnnlab_lang", "vi")
    exact = _UI_TRANSLATIONS.get(value)
    if exact:
        return exact.get(lang, value)
    # Also translate known phrases inside dynamic messages (e.g. f-strings).
    # Single-pass replacement avoids translating a phrase's newly inserted
    # target text a second time (which can otherwise mangle capitalization).
    keys = sorted(_UI_TRANSLATIONS, key=len, reverse=True)
    alternatives = []
    for source in keys:
        if len(source) < 4:
            alternatives.append(r"(?<!\w)" + re.escape(source) + r"(?!\w)")
        else:
            alternatives.append(re.escape(source))
    pattern = re.compile("|".join(alternatives))
    return pattern.sub(
        lambda match: _UI_TRANSLATIONS[match.group(0)].get(lang, match.group(0)),
        value,
    )


def _translate_nested(value):
    if isinstance(value, dict):
        return {translate_ui(k) if isinstance(k, str) else k: _translate_nested(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_translate_nested(v) for v in value]
    if isinstance(value, tuple):
        return tuple(_translate_nested(v) for v in value)
    if isinstance(value, str):
        return translate_ui(value)
    return value


def install_streamlit_i18n():
    """Localize Streamlit display text, widgets, tables, and JSON in every container."""
    try:
        import inspect
        from streamlit.delta_generator import DeltaGenerator
    except Exception:
        return
    if getattr(DeltaGenerator, "_cnnlab_i18n_installed", False):
        return

    display_methods = {
        "title", "header", "subheader", "caption", "markdown", "text",
        "info", "success", "warning", "error", "exception", "toast", "write",
    }
    labeled_methods = {
        "button", "download_button", "checkbox", "radio", "selectbox", "multiselect",
        "text_input", "text_area", "number_input", "slider", "select_slider",
        "file_uploader", "form_submit_button", "metric", "expander", "spinner",
        "status", "progress", "tabs", "form",
    }
    method_names = display_methods | labeled_methods | {"dataframe", "table", "json", "image"}

    def make_wrapper(method_name, original):
        try:
            signature = inspect.signature(original)
        except (TypeError, ValueError):
            signature = None

        def wrapped(self, *args, **kwargs):
            if signature is None:
                return original(self, *args, **kwargs)
            try:
                bound = signature.bind_partial(self, *args, **kwargs)
                values = bound.arguments

                if method_name in display_methods:
                    # Most text display methods call their content `body` or `text`;
                    # write() accepts a variadic tuple.
                    for key in ("body", "text", "message", "args"):
                        if key in values:
                            values[key] = _translate_nested(values[key])
                elif method_name in labeled_methods:
                    for key in ("label", "help", "placeholder", "caption", "text"):
                        if key in values and isinstance(values[key], str):
                            values[key] = translate_ui(values[key])
                    if method_name == "metric" and isinstance(values.get("value"), str):
                        values["value"] = translate_ui(values["value"])
                    # Translate option presentation only. The selected values remain
                    # the original strings, so model configuration and branching work.
                    if method_name in {"radio", "selectbox", "multiselect", "select_slider"} and "options" in values:
                        old_format = values.get("format_func")
                        if old_format is None:
                            old_format = lambda x: x
                        values["format_func"] = lambda x, _old=old_format: translate_ui(_old(x))
                elif method_name in {"dataframe", "table"}:
                    data = values.get("data")
                    if isinstance(data, pd.DataFrame):
                        data = data.copy()
                        class_columns = {"actual", "y_true", "predicted_class", "prediction", "class_name", "true_label", "predicted_label"}
                        for col in data.columns:
                            if str(col).lower() in class_columns:
                                data[col] = data[col].map(
                                    lambda v: translate_ui(v) if isinstance(v, str) and v in {"Cat", "Dog"} else v
                                )
                        data.rename(columns={col: translate_ui(str(col)) for col in data.columns}, inplace=True)
                        if data.index.name:
                            data.index.name = translate_ui(str(data.index.name))
                        values["data"] = data
                elif method_name == "json":
                    for key in ("body",):
                        if key in values:
                            values[key] = _translate_nested(values[key])
                elif method_name == "image":
                    if "caption" in values:
                        values["caption"] = _translate_nested(values["caption"])
            except Exception:
                # If argument inspection/localization fails, make one normal call.
                return original(self, *args, **kwargs)
            # Keep the original call outside the fallback handler so an actual
            # Streamlit exception is not accidentally triggered a second time.
            return original(*bound.args, **bound.kwargs)

        wrapped.__name__ = getattr(original, "__name__", method_name)
        wrapped.__doc__ = getattr(original, "__doc__", None)
        return wrapped

    for name in method_names:
        original = getattr(DeltaGenerator, name, None)
        if callable(original):
            setattr(DeltaGenerator, name, make_wrapper(name, original))
    setattr(DeltaGenerator, "_cnnlab_i18n_installed", True)

def build_survey_docx(selected_runs, language):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    labels = I18N[language]
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65); sec.bottom_margin = Inches(0.65)
    styles = doc.styles
    styles['Normal'].font.name = 'Arial'; styles['Normal'].font.size = Pt(10)
    title = doc.add_heading(labels['report_title'], 0); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"CS535 | CNN Research Lab | {translate_ui('Generated:')} {time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_heading(labels['procedure'], level=1)
    doc.add_paragraph('Dataset QC → saved stratified split (70/15/15) → train runs → choose using validation metrics → lock final model → independent test → error analysis.' if language=='en' else ('資料集品質檢查 → 固定分層切分（70/15/15）→ 訓練實驗 → 依驗證指標選擇 → 鎖定最終模型 → 獨立測試 → 錯誤分析。' if language=='zh-TW' else 'Kiểm tra chất lượng dữ liệu → split phân tầng đã lưu (70/15/15) → huấn luyện → chọn bằng validation → khóa mô hình cuối → independent test → phân tích lỗi.'))
    doc.add_paragraph(translate_ui('Important:') + (' independent-test results must not be used to select hyperparameters or compare candidate models.' if language=='en' else ('不得使用獨立測試集結果選擇超參數或比較候選模型。' if language=='zh-TW' else ' không dùng kết quả independent test để chọn siêu tham số hoặc so sánh mô hình ứng viên.'))) 
    doc.add_heading(labels['results'], level=1)
    table=doc.add_table(rows=1, cols=6); table.style='Light Shading Accent 1'
    heads=[translate_ui('Run ID'),translate_ui('Model name'),translate_ui('created_at'),translate_ui('val_accuracy'),translate_ui('val_f1'),translate_ui('val_roc_auc')]
    for cell, txt in zip(table.rows[0].cells, heads): cell.text=txt
    for r in selected_runs:
        row=table.add_row().cells
        for cell, val in zip(row,[r.get('run_id','—'),r.get('model_name','—'),r.get('created_at','—'),r.get('val_accuracy','—'),r.get('val_f1','—'),r.get('val_roc_auc','—')]): cell.text=str(val if val is not None else '—')
    for idx, r in enumerate(selected_runs, 1):
        run_dir=Path(r['run_dir']); cfg=read_json(run_dir/'config.json'); metrics=read_json(run_dir/'validation_metrics.json')
        hist_path=run_dir/'history.csv'; hist=pd.read_csv(hist_path) if hist_path.exists() else pd.DataFrame()
        doc.add_heading(f"{idx}. {r.get('run_id','Run')} — {r.get('model_name','CNN')}", level=1)
        doc.add_heading(labels['config'], level=2)
        tbl=doc.add_table(rows=0, cols=2); tbl.style='Light List Accent 1'
        for k,v in cfg.items():
            if k in ('dataset_manifest','split_manifest_sha256'): continue
            cells=tbl.add_row().cells; cells[0].text=translate_ui(str(k)); cells[1].text=json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else str(v)
        doc.add_heading(labels['results'], level=2)
        mt=doc.add_table(rows=1, cols=2); mt.style='Light Shading Accent 1'; mt.rows[0].cells[0].text=labels['metric']; mt.rows[0].cells[1].text=labels['value']
        for k,v in metrics.items():
            if isinstance(v,(int,float,str)) or v is None:
                cells=mt.add_row().cells; cells[0].text=translate_ui(str(k)); cells[1].text=str(v if v is not None else '—')
        if not hist.empty:
            doc.add_heading(translate_ui('Learning history'), level=2)
            cols=[c for c in ['epoch','train_loss','val_loss','val_accuracy','val_roc_auc','val_f1'] if c in hist.columns]
            ht=doc.add_table(rows=1, cols=len(cols)); ht.style='Light List Accent 1'
            for c, name in zip(ht.rows[0].cells,cols): c.text=translate_ui(name)
            for _, hr in hist[cols].iterrows():
                cells=ht.add_row().cells
                for c, name in zip(cells,cols): c.text=str(round(float(hr[name]),5)) if pd.notna(hr[name]) and isinstance(hr[name],(int,float)) else str(hr[name])
        doc.add_heading(labels['interpretation'], level=2)
        notes=[]
        acc=metrics.get('accuracy'); f1=metrics.get('f1'); auc=metrics.get('roc_auc')
        if acc is not None: notes.append(f"{translate_ui('Validation accuracy =')} {float(acc):.4f}.")
        if f1 is not None: notes.append(f"{translate_ui('Validation F1 =')} {float(f1):.4f}.")
        if auc is not None: notes.append(f"{translate_ui('Validation ROC-AUC =')} {float(auc):.4f}.")
        if len(hist)>=2 and 'train_loss' in hist and 'val_loss' in hist:
            first,last=hist.iloc[0],hist.iloc[-1]
            if pd.notna(first['train_loss']) and pd.notna(last['train_loss']):
                notes.append(f"{translate_ui('Training loss changed from')} {first['train_loss']:.4f} {'đến' if language=='vi' else ('至' if language=='zh-TW' else 'to')} {last['train_loss']:.4f}.")
            if pd.notna(last.get('train_loss')) and pd.notna(last.get('val_loss')) and last['val_loss']>last['train_loss']*1.25:
                notes.append(translate_ui('Final validation loss is notably higher than training loss; inspect possible overfitting.'))
            elif pd.notna(last.get('train_loss')) and pd.notna(last.get('val_loss')):
                notes.append(translate_ui('Compare the training and validation curves to assess generalization.'))
        doc.add_paragraph(' '.join(notes) if notes else translate_ui('Insufficient saved metrics for an evidence-based observation.'))
    doc.add_heading(labels['limitations'], level=1)
    doc.add_paragraph('This report summarizes saved run metadata and metrics. It does not establish statistical significance, and it does not claim that any model is best unless a controlled validation comparison supports that conclusion. Independent-test metrics are included only if already saved.' if language=='en' else ('本報告彙整已保存的實驗設定與指標，不代表統計顯著性；只有在受控驗證比較支持時，才能判定模型較佳。獨立測試指標僅在已保存時列入。' if language=='zh-TW' else 'Báo cáo tổng hợp cấu hình và metric đã lưu, không khẳng định ý nghĩa thống kê hay mô hình tốt nhất nếu chưa có so sánh validation có kiểm soát. Chỉ đưa metric independent test nếu đã được lưu.'))
    out=BytesIO(); doc.save(out); return out.getvalue()



def read_json(path: Path):
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def metric_value(value, digits=3):
    try:
        if value is None or pd.isna(value):
            return "—"
        if isinstance(value, float):
            return f"{value:.{digits}f}"
        return str(value)
    except Exception:
        return str(value)


def get_split_path() -> Path:
    return processed_dir() / "split_manifest.csv"


def ensure_split_df() -> pd.DataFrame | None:
    path = get_split_path()
    if path.exists():
        try:
            df = pd.read_csv(path)
            if {"path", "label", "split"}.issubset(df.columns):
                return df
        except Exception:
            return None
    return None


def experiment_choices(rows):
    return {
        f"{r.get('run_id', 'unknown')} · {r.get('model_name', 'CNN')} · val AUC {metric_value(r.get('val_roc_auc'))}": r
        for r in rows
    }


# Read the selected language before rendering the page, so every screen updates
# in the same rerun. A stable key preserves the choice between reruns.
st.session_state.setdefault("cnnlab_language_picker", "Tiếng Việt")
st.session_state["cnnlab_lang"] = LANGUAGES.get(st.session_state["cnnlab_language_picker"], "vi")
lang_name = st.sidebar.selectbox(
    tr("language"), list(LANGUAGES.keys()),
    index=list(LANGUAGES.keys()).index(st.session_state["cnnlab_language_picker"]),
    key="cnnlab_language_picker",
)
st.session_state["cnnlab_lang"] = LANGUAGES[lang_name]
install_streamlit_i18n()

st.title(tr("title"))
st.caption(tr("caption"))
st.markdown(
    "Quản lý dữ liệu, huấn luyện CNN, so sánh thí nghiệm và phân tích lỗi trong cùng một giao diện. "
    "Ứng dụng giữ riêng validation và independent test để hỗ trợ quy trình nghiên cứu có thể tái lập."
)

with st.sidebar:
    st.caption(tr("language_help"))
    st.header(tr("lab"))
    page_keys = ["overview", "dataset", "training", "experiments", "final", "errors", "survey"]
    page = st.radio("Chọn phân hệ", [tr(k) for k in page_keys], label_visibility="collapsed")
    page = dict(zip([tr(k) for k in page_keys], page_keys))[page]
    st.divider()
    st.caption(f"Thiết bị: {'CUDA GPU' if torch.cuda.is_available() else 'CPU'}")
    st.caption(f"Seed mặc định: {SEED}")
    st.caption("Thư mục dự án")
    st.code(str(project_root()), language=None)

split_df = ensure_split_df()
experiments = list_experiments()

if page == "overview":
    st.subheader("Bảng điều khiển nghiên cứu")
    n_valid = len(split_df) if split_df is not None else 0
    n_runs = len(experiments)
    test_runs = sum(1 for r in experiments if (Path(r["run_dir"]) / "internal_test_metrics.json").exists())
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Ảnh hợp lệ trong split", f"{n_valid:,}" if n_valid else "Chưa chuẩn bị")
    c2.metric("Thí nghiệm đã lưu", n_runs)
    c3.metric("Run đã đánh giá test", test_runs)
    c4.metric("Thiết bị", "GPU" if torch.cuda.is_available() else "CPU")

    st.subheader("Quy trình thực nghiệm")
    st.markdown(
        "**Dataset → QC/manifest → stratified split 70/15/15 → baseline/controlled runs → "
        "chọn cấu hình bằng validation → independent test một lần sau khi khóa mô hình → error analysis.**"
    )
    st.warning(
        "Không chọn mô hình hoặc chỉnh siêu tham số dựa trên independent test. "
        "Trong giai đoạn phát triển, hãy dùng validation metrics và validation errors."
    )
    st.subheader("Các bước bắt đầu")
    st.markdown(
        "1. Mở **Dataset & QC** để tải dataset Microsoft hoặc quét thư mục ảnh sẵn có.\n"
        "2. Tạo manifest và split tái lập.\n"
        "3. Mở **Huấn luyện** để chạy baseline CNN hoặc kiến trúc TorchVision.\n"
        "4. So sánh các run tại **Thí nghiệm**; khi đã khóa lựa chọn, mới mở **Đánh giá cuối & Challenge**."
    )
    if experiments:
        st.subheader("Các run gần đây")
        view = pd.DataFrame(experiments)
        cols = [c for c in ["run_id", "model_name", "created_at", "train_images_used", "val_accuracy", "val_roc_auc", "val_f1", "elapsed_seconds"] if c in view]
        st.dataframe(view[cols].head(8), use_container_width=True, hide_index=True)
    else:
        st.info("Chưa có run nào. Hãy chuẩn bị dataset trước khi huấn luyện.")

elif page == "dataset":
    st.subheader("Quản lý dataset và kiểm tra chất lượng")
    st.write("Nguồn chính được khuyến nghị trong tài liệu CS535 là Microsoft/Asirra Cats vs. Dogs (~25.000 ảnh thô).")
    st.markdown("Nguồn chính thức: [Microsoft Cats and Dogs](https://www.microsoft.com/en-us/download/details.aspx?id=54765)")
    raw_default = str(data_root() / "raw")
    root = st.text_input("Thư mục dataset", value=raw_default, help="Có thể là thư mục chứa PetImages, hoặc chính thư mục có hai thư mục Cat và Dog.")
    left, right = st.columns(2)
    with left:
        if st.button("Tải dataset Microsoft (~786 MB)", type="primary", use_container_width=True):
            progress = st.progress(0.0)
            status = st.empty()
            def dl_cb(frac, message):
                progress.progress(max(0.0, min(1.0, float(frac))))
                status.write(message)
            try:
                safe_download_microsoft_dataset(Path(root), dl_cb)
                st.success(f"Đã tải/giải nén dữ liệu trong {root}")
            except Exception as exc:
                st.error(f"Không thể tải hoặc giải nén: {type(exc).__name__}: {exc}")
    with right:
        if st.button("Quét ảnh, loại file lỗi/trùng và tạo split", use_container_width=True):
            status = st.empty()
            status.info("Đang quét ảnh và kiểm tra file. Với khoảng 25.000 ảnh, bước này có thể mất vài phút.")
            try:
                manifest, report = scan_dataset(root)
                splits = make_splits(manifest, seed=SEED)
                split_df = splits
                target = processed_dir()
                split_path = target / "split_manifest.csv"
                old_fingerprint = hashlib.sha256(split_path.read_bytes()).hexdigest() if split_path.exists() else None
                new_split_bytes = splits.to_csv(index=False).encode("utf-8")
                new_fingerprint = hashlib.sha256(new_split_bytes).hexdigest()
                if old_fingerprint and old_fingerprint != new_fingerprint:
                    lock_path = target / "final_test_lock.json"
                    if lock_path.exists():
                        archive_path = target / f"final_test_lock_archived_{time.strftime('%Y%m%d_%H%M%S')}.json"
                        lock_path.replace(archive_path)
                manifest.to_csv(target / "manifest.csv", index=False)
                split_path.write_bytes(new_split_bytes)
                (target / "dataset_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
                st.session_state["last_scan_report"] = report
                status.success("Quét và chia tập dữ liệu hoàn tất.")
            except Exception as exc:
                status.empty()
                st.error(f"Không thể chuẩn bị dataset: {type(exc).__name__}: {exc}")

    report = read_json(processed_dir() / "dataset_report.json")
    if report:
        st.subheader("Báo cáo kiểm tra dữ liệu gần nhất")
        a, b, c, d = st.columns(4)
        a.metric("Ảnh hợp lệ sau dedup", f"{report.get('valid_images_after_dedup', 0):,}")
        b.metric("Cat", f"{report.get('valid_cat', 0):,}")
        c.metric("Dog", f"{report.get('valid_dog', 0):,}")
        d.metric("File lỗi / bỏ qua", f"{report.get('invalid_files', 0):,}")
        st.caption(f"Exact duplicate đã loại: {report.get('exact_duplicates_removed', 0):,} · Trùng ảnh nhưng mâu thuẫn nhãn đã loại: {report.get('cross_class_conflicts_removed', 0):,} · Quét lúc: {report.get('scanned_at', '—')}")
        if report.get("invalid_examples"):
            with st.expander("Một số file không hợp lệ (tối đa 100 mục được lưu)"):
                st.dataframe(pd.DataFrame(report["invalid_examples"]), use_container_width=True, hide_index=True)
    if split_df is not None:
        st.subheader("Phân chia train / validation / independent test")
        summary = split_df.groupby(["split", "class_name"]).size().unstack(fill_value=0)
        summary["Total"] = summary.sum(axis=1)
        st.dataframe(summary, use_container_width=True)
        st.caption("Seed = 42; stratified 70/15/15. Split đã được lưu để các thí nghiệm dùng cùng một phân chia.")
        if "width" in split_df and "height" in split_df:
            with st.expander("Thống kê kích thước ảnh nguồn"):
                st.dataframe(split_df[["width", "height", "min_side", "max_side", "aspect_ratio"]].describe().round(2), use_container_width=True)
    else:
        st.info("Chưa có split manifest. Quét dataset để hệ thống tạo manifest và chia tập.")
    with st.expander("Dataset đánh giá ngoài: Oxford-IIIT Pet"):
        st.write("Theo đề xuất nghiên cứu, có thể giữ Oxford-IIIT Pet làm external test độc lập. MVP này không tự tải/chuyển nhãn Oxford; hãy chuẩn bị thư mục chuẩn hóa cat/ và dog/ để đánh giá ở trang Challenge.")
        st.markdown("[Oxford-IIIT Pet dataset](https://www.robots.ox.ac.uk/~vgg/data/pets/)")

elif page == "training":
    st.subheader("Cấu hình một thí nghiệm")
    if split_df is None:
        st.error("Chưa có dataset split. Vào Dataset & QC để quét dữ liệu trước.")
        st.stop()
    with st.form("experiment_config"):
        model_name = st.selectbox("Kiến trúc", ["configurable_cnn", "resnet18", "mobilenet_v3_small", "efficientnet_b0", "convnext_tiny"],
                                  format_func=lambda x: {"configurable_cnn":"Configurable CNN (from scratch)", "resnet18":"ResNet-18", "mobilenet_v3_small":"MobileNetV3-Small", "efficientnet_b0":"EfficientNet-B0", "convnext_tiny":"ConvNeXt-Tiny"}[x])
        c1, c2, c3 = st.columns(3)
        with c1:
            image_size = st.selectbox("Image resolution", [64, 96, 128, 160, 224], index=2)
            batch_size = st.selectbox("Batch size", [8, 16, 32, 64], index=2)
            epochs = st.slider("Epochs", min_value=1, max_value=30, value=3)
        with c2:
            optimizer = st.selectbox("Optimizer", ["AdamW", "Adam", "SGD"])
            learning_rate = st.select_slider("Learning rate", options=[0.0001, 0.0005, 0.001, 0.01], value=0.001, format_func=lambda x: f"{x:g}")
            weight_decay = st.select_slider("Weight decay", options=[0.0, 0.00001, 0.0001, 0.001], value=0.0001, format_func=lambda x: f"{x:g}")
        with c3:
            max_train_images = st.selectbox("Train images per run", [1000, 2500, 5000, 10000, 0], index=2, format_func=lambda x: "Toàn bộ train split" if x == 0 else f"Tối đa {x:,} (quick run)")
            augmentation = st.checkbox("Data augmentation nhẹ", value=True)
            transfer_strategy = st.selectbox(
                "Transfer-learning strategy",
                ["scratch", "feature_extraction", "fine_tuning"],
                format_func=lambda x: {"scratch": "From scratch", "feature_extraction": "Feature extraction (freeze backbone)", "fine_tuning": "Fine-tuning (pretrained backbone)"}[x],
                disabled=(model_name == "configurable_cnn"),
                help="Dùng với kiến trúc TorchVision. Feature extraction chỉ cập nhật classifier; fine-tuning cập nhật toàn bộ model.",
            )
        if model_name == "configurable_cnn":
            transfer_strategy = "scratch"
        pretrained = transfer_strategy != "scratch" and model_name != "configurable_cnn"
        if transfer_strategy == "fine_tuning":
            st.caption("Fine-tuning thường nên bắt đầu với learning rate nhỏ (ví dụ 1e-4) để tránh phá vỡ trọng số pretrained.")
        st.caption("Pretrained weights có thể cần tải từ Internet ở lần chạy đầu. Mọi kiến trúc đều dùng cùng split đã lưu.")
        st.markdown("**Tham số riêng của Configurable CNN**")
        c4, c5, c6 = st.columns(3)
        with c4:
            depth = st.select_slider("Depth / convolution blocks", options=[2, 3, 4, 5], value=3)
            base_filters = st.selectbox("Base filters", [16, 32, 64, 96, 128], index=1)
        with c5:
            kernel_size = st.selectbox("Kernel size", [1, 2, 3, 4, 5, 6, 7, 10], index=2)
            pool = st.selectbox("Pooling", ["max", "avg", "none"], index=0)
        with c6:
            dropout = st.select_slider("Dropout", options=[0.0, 0.2, 0.3, 0.5], value=0.3)
            batchnorm = st.checkbox("Batch normalization", value=True)
            activation = st.selectbox("Activation", ["relu", "silu"])
        seed = st.number_input("Random seed", min_value=0, max_value=999999, value=SEED, step=1)
        submitted = st.form_submit_button("Bắt đầu huấn luyện", type="primary", use_container_width=True)

    if submitted:
        cfg = {
            "model_name": model_name, "pretrained": bool(pretrained), "transfer_strategy": transfer_strategy,
            "image_size": int(image_size), "batch_size": int(batch_size), "epochs": int(epochs),
            "optimizer": optimizer, "learning_rate": float(learning_rate), "weight_decay": float(weight_decay),
            "max_train_images": int(max_train_images), "augmentation": bool(augmentation),
            "depth": int(depth), "base_filters": int(base_filters), "kernel_size": int(kernel_size),
            "pool": pool, "dropout": float(dropout), "batchnorm": bool(batchnorm),
            "activation": activation, "seed": int(seed),
            "split_ratio": {"train": 0.70, "validation": 0.15, "test": 0.15},
            "dataset_manifest": str(get_split_path()),
        }
        progress = st.progress(0.0)
        status = st.empty()
        live = st.empty()
        def on_progress(item):
            progress.progress(min(item["epoch"] / item["epochs"], 1.0))
            status.write(f"Epoch {item['epoch']}/{item['epochs']} · loss train={item['train_loss']:.4f} · loss val={item['val_loss']:.4f}")
            live.write(f"Validation accuracy: {item['val_accuracy']:.3f} · Validation ROC-AUC: {metric_value(item['val_roc_auc'])} · thời gian epoch: {item['epoch_seconds']:.1f}s")
        try:
            with st.spinner("Đang khởi tạo model và huấn luyện..."):
                run_dir = train_experiment(split_df, cfg, on_progress)
            st.success(f"Đã lưu thí nghiệm: {run_dir.name}")
            st.write("Checkpoint, config, history, validation metrics và validation predictions đã được lưu trong thư mục run.")
            metrics = read_json(run_dir / "validation_metrics.json")
            m1, m2, m3, m4 = st.columns(4)
            m1.metric("Validation accuracy", metric_value(metrics.get("accuracy")))
            m2.metric("Validation ROC-AUC", metric_value(metrics.get("roc_auc")))
            m3.metric("Validation F1", metric_value(metrics.get("f1")))
            m4.metric("Trainable parameters", f"{read_json(run_dir / 'config.json').get('parameters', 0):,}")
            st.info("Các chỉ số đang hiển thị là validation metrics; internal test chưa được sử dụng.")
        except Exception as exc:
            st.error(f"Huấn luyện thất bại: {type(exc).__name__}: {exc}")

elif page == "experiments":
    st.subheader("Experiment tracking & comparison")
    if not experiments:
        st.info("Chưa có thí nghiệm. Hãy chạy ít nhất một baseline.")
    else:
        df = pd.DataFrame(experiments)
        if "split_manifest_sha256" in df:
            df["split_id"] = df["split_manifest_sha256"].fillna("").astype(str).str[:10]
        view_cols = [c for c in ["run_id", "model_name", "created_at", "split_id", "image_size", "epochs", "train_images_used", "parameters", "val_accuracy", "val_balanced_accuracy", "val_precision", "val_recall_sensitivity", "val_specificity", "val_f1", "val_roc_auc", "val_pr_auc_avg_precision", "elapsed_seconds"] if c in df]
        st.dataframe(df[view_cols], use_container_width=True, hide_index=True)
        st.download_button("Tải bảng so sánh CSV", df[view_cols].to_csv(index=False).encode("utf-8-sig"), file_name="cnn_experiment_comparison.csv", mime="text/csv")
        st.subheader("So sánh validation metrics")
        metric = st.selectbox("Chỉ số", ["val_roc_auc", "val_accuracy", "val_f1", "val_balanced_accuracy", "val_pr_auc_avg_precision"], index=0)
        plot_df = df.dropna(subset=[metric]).copy() if metric in df else pd.DataFrame()
        if not plot_df.empty:
            plot_df["label"] = plot_df["run_id"].astype(str) + " · " + plot_df["model_name"].astype(str)
            fig, ax = plt.subplots(figsize=(10, max(3.5, 0.42 * len(plot_df))))
            ordered = plot_df.sort_values(metric)
            ax.barh(ordered["label"], ordered[metric])
            ax.set_xlabel(translate_ui(metric))
            ax.grid(axis="x", alpha=0.25)
            fig.tight_layout()
            st.pyplot(fig)
            plt.close(fig)
        st.subheader("Learning curve")
        selected_label = st.selectbox("Chọn run để xem learning curve", list(experiment_choices(experiments).keys()))
        selected = experiment_choices(experiments)[selected_label]
        history_path = Path(selected["run_dir"]) / "history.csv"
        if history_path.exists():
            hist = pd.read_csv(history_path)
            fig, ax = plt.subplots(figsize=(9, 3.7))
            ax.plot(hist["epoch"], hist["train_loss"], marker="o", label=translate_ui("Train loss"))
            ax.plot(hist["epoch"], hist["val_loss"], marker="o", label=translate_ui("Validation loss"))
            ax.set_xlabel(translate_ui("Epoch")); ax.set_ylabel(translate_ui("BCE loss")); ax.legend(); ax.grid(alpha=0.25)
            fig.tight_layout(); st.pyplot(fig); plt.close(fig)
            fig, ax = plt.subplots(figsize=(9, 3.7))
            ax.plot(hist["epoch"], hist["val_accuracy"], marker="o", label=translate_ui("Validation accuracy"))
            ax.plot(hist["epoch"], hist["val_roc_auc"], marker="o", label=translate_ui("Validation ROC-AUC"))
            ax.set_xlabel(translate_ui("Epoch")); ax.set_ylim(0, 1); ax.legend(); ax.grid(alpha=0.25)
            fig.tight_layout(); st.pyplot(fig); plt.close(fig)
        with st.expander("Chi tiết cấu hình của run đã chọn"):
            st.json(read_json(Path(selected["run_dir"]) / "config.json"))
        st.warning("Sử dụng validation để so sánh. Không xếp hạng/chọn run theo internal test metrics.")

elif page == "final":
    st.subheader("Independent test và external challenge")
    st.warning(
        "Chỉ chạy internal test sau khi đã khóa mô hình và siêu tham số. Không dùng kết quả này để quay lại chọn run; "
        "mỗi lần đánh giá test đều nên được ghi nhận trong báo cáo nghiên cứu."
    )
    if not experiments:
        st.info("Chưa có checkpoint để đánh giá.")
    else:
        choices = experiment_choices(experiments)
        label = st.selectbox("Chọn checkpoint đã chốt để đánh giá", list(choices.keys()), key="final_eval_run")
        selected = choices[label]
        run_dir = Path(selected["run_dir"])
        lock_path = processed_dir() / "final_test_lock.json"
        test_lock = read_json(lock_path)
        locked_run_id = test_lock.get("run_id")
        selected_config = read_json(run_dir / "config.json")
        split_path = get_split_path()
        current_split_fingerprint = hashlib.sha256(split_path.read_bytes()).hexdigest() if split_path.exists() else None
        same_split = bool(current_split_fingerprint and selected_config.get("split_manifest_sha256") == current_split_fingerprint)
        if not same_split:
            st.error("Run này không được huấn luyện trên split manifest hiện tại. Không thể dùng internal test hiện tại để đánh giá run này một cách hợp lệ.")
        if test_lock and locked_run_id != selected.get("run_id"):
            st.error(
                f"Independent test đã được sử dụng cho run {locked_run_id}. "
                "Theo quy trình nghiên cứu, không chạy lại test trên các run khác để chọn mô hình."
            )
        if (run_dir / "internal_test_metrics.json").exists():
            st.info("Run này có kết quả internal test đã lưu. Kết quả này chỉ dành cho báo cáo cuối.")
            metrics = read_json(run_dir / "internal_test_metrics.json")
            show = ["accuracy", "balanced_accuracy", "precision", "recall_sensitivity", "specificity", "f1", "roc_auc", "pr_auc_avg_precision", "mcc", "brier"]
            cols = st.columns(5)
            for i, key in enumerate(show):
                cols[i % 5].metric(key, metric_value(metrics.get(key)))
        test_already_done = bool(test_lock) or (run_dir / "internal_test_metrics.json").exists()
        if st.button("Independent test đã khóa" if test_already_done else "Chạy internal test cho checkpoint đã chốt", type="primary", disabled=(test_already_done or not same_split)):
            if split_df is None:
                st.error("Không tìm thấy split manifest.")
            else:
                try:
                    with st.spinner("Đang chạy independent test..."):
                        metrics, preds = evaluate_internal_test(run_dir, get_split_path())
                    lock_path.write_text(json.dumps({
                        "run_id": selected.get("run_id"),
                        "evaluated_at": metrics.get("evaluated_at"),
                        "split_manifest_sha256": read_json(Path(selected["run_dir"]) / "config.json").get("split_manifest_sha256"),
                        "policy": "one final model per saved split; do not compare test results across runs",
                    }, ensure_ascii=False, indent=2), encoding="utf-8")
                    st.success(f"Đã lưu đánh giá independent test với n={metrics.get('n_samples', 0)} ảnh và khóa run cuối.")
                    st.json(metrics)
                    st.dataframe(preds.head(30), use_container_width=True, hide_index=True)
                except Exception as exc:
                    st.error(f"Test evaluation thất bại: {type(exc).__name__}: {exc}")
        st.divider()
        st.subheader("External / hard challenge set")
        st.write("Chuẩn bị thư mục ngoài dataset theo cấu trúc: `challenge_set/cat/`, `challenge_set/dog/`, `challenge_set/ood/`. Nhãn cat/dog dùng để tính metric; OOD chỉ phân tích confidence, không tính binary accuracy.")
        challenge_root = st.text_input("Đường dẫn challenge_set", value=str(project_root() / "challenge_set"), key="challenge_root")
        challenge_allowed = bool(test_lock and locked_run_id == selected.get("run_id") and same_split)
        if not challenge_allowed:
            st.info("Chỉ đánh giá external challenge sau khi đã khóa mô hình cuối qua internal test. Dùng validation để chọn mô hình trước bước đó.")
        if st.button("Đánh giá challenge folder", disabled=not challenge_allowed):
            try:
                metrics, preds = evaluate_challenge_folder(run_dir, challenge_root)
                st.success(f"Đã đánh giá {metrics.get('n_images', 0)} ảnh challenge; không cập nhật trọng số mô hình.")
                st.json(metrics)
                st.dataframe(preds.head(50), use_container_width=True, hide_index=True)
            except Exception as exc:
                st.error(f"Challenge evaluation thất bại: {type(exc).__name__}: {exc}")
        if (run_dir / "challenge_metrics.json").exists():
            with st.expander("Kết quả challenge đã lưu"):
                st.json(read_json(run_dir / "challenge_metrics.json"))

elif page == "errors":
    st.subheader("Phân tích lỗi và dự đoán ảnh mới")
    if not experiments:
        st.info("Chưa có checkpoint. Huấn luyện một run trước.")
        st.stop()
    choices = experiment_choices(experiments)
    label = st.selectbox("Checkpoint", list(choices.keys()), key="analysis_run")
    selected = choices[label]
    run_dir = Path(selected["run_dir"])
    source_options = {"Validation (khuyến nghị trong phát triển)": "validation_predictions.csv"}
    if (run_dir / "internal_test_predictions.csv").exists():
        source_options["Internal test (chỉ xem báo cáo cuối) "] = "internal_test_predictions.csv"
    if (run_dir / "challenge_predictions.csv").exists():
        source_options["External / hard challenge"] = "challenge_predictions.csv"
    source_label = st.selectbox("Nguồn phân tích lỗi", list(source_options.keys()))
    preds_path = run_dir / source_options[source_label]
    if preds_path.exists():
        preds = pd.read_csv(preds_path)
        if "correct" not in preds.columns and {"y_true", "prediction"}.issubset(preds.columns):
            preds["correct"] = (preds["y_true"] == preds["prediction"]).where(preds["y_true"].notna())
        if "confidence_in_prediction" not in preds.columns and "prob_dog" in preds:
            preds["confidence_in_prediction"] = preds["prob_dog"].apply(lambda p: max(float(p), 1 - float(p)))
        wrong = preds[preds["correct"] == False].copy() if "correct" in preds else pd.DataFrame()
        c1, c2, c3 = st.columns(3)
        c1.metric("Số ảnh trong file", len(preds))
        c2.metric("Lỗi dự đoán", len(wrong))
        c3.metric("Tỷ lệ lỗi", metric_value(len(wrong) / len(preds) if len(preds) else None))
        if not wrong.empty:
            st.markdown("**Các trường hợp dự đoán sai có confidence cao**")
            wrong = wrong.sort_values("confidence_in_prediction", ascending=False)
            st.dataframe(wrong.head(30), use_container_width=True, hide_index=True)
            show_n = st.slider("Số ảnh lỗi để xem", min_value=1, max_value=min(12, len(wrong)), value=min(8, len(wrong)))
            cols = st.columns(4)
            for i, row in enumerate(wrong.head(show_n).itertuples(index=False)):
                path = getattr(row, "path", "")
                with cols[i % 4]:
                    try:
                        st.image(Image.open(path).convert("RGB"), caption=f"True: {getattr(row, 'actual', getattr(row, 'y_true', '?'))} · Pred: {getattr(row, 'predicted_class', getattr(row, 'prediction', '?'))} · p(Dog)={float(getattr(row, 'prob_dog', 0)):.2f}", use_container_width=True)
                    except Exception:
                        st.caption(f"Không mở được ảnh: {path}")
        else:
            st.success("Không tìm thấy trường hợp dự đoán sai trong file đã chọn, hoặc nguồn dữ liệu không có nhãn.")
        st.download_button("Tải prediction table CSV", preds.to_csv(index=False).encode("utf-8-sig"), file_name=f"{run_dir.name}_{preds_path.stem}.csv", mime="text/csv")
    else:
        st.info("Chưa có prediction file cho nguồn này.")

    st.divider()
    st.subheader("Thử dự đoán một ảnh mới")
    uploaded = st.file_uploader("Tải ảnh JPG / PNG / WEBP", type=["jpg", "jpeg", "png", "webp"])
    if uploaded is not None:
        image = Image.open(uploaded).convert("RGB")
        st.image(image, caption="Ảnh đầu vào", width=320)
        if st.button("Phân loại ảnh", type="primary"):
            try:
                model, cfg, device = load_experiment(run_dir)
                tensor = make_transform(int(cfg.get("image_size", 128)), train=False, augment=False)(image).unsqueeze(0).to(device)
                with torch.no_grad():
                    probability_dog = float(torch.sigmoid(model(tensor).reshape(-1))[0].item())
                prediction = "Dog" if probability_dog >= 0.5 else "Cat"
                st.metric("Dự đoán", prediction)
                st.write(f"P(Dog) = **{probability_dog:.4f}** · P(Cat) = **{1 - probability_dog:.4f}**")
                st.caption("Xác suất là đầu ra sigmoid của bộ phân loại nhị phân; không tự động đồng nghĩa với xác suất đã hiệu chuẩn.")
            except Exception as exc:
                st.error(f"Không thể dự đoán: {type(exc).__name__}: {exc}")


elif page == "survey":
    st.subheader(tr("report_title"))
    st.write(tr("report_intro"))
    st.caption(tr("report_note"))
    if not experiments:
        st.info(tr("no_runs"))
    else:
        choices = experiment_choices(experiments)
        labels = list(choices.keys())
        selected_labels = st.multiselect(tr("select_runs"), labels, default=labels)
        if st.button(tr("generate"), type="primary", disabled=not selected_labels):
            chosen = [choices[label] for label in selected_labels]
            try:
                content = build_survey_docx(chosen, st.session_state.get("cnnlab_lang", "vi"))
                st.session_state["cnnlab_survey_docx"] = content
                st.session_state["cnnlab_survey_name"] = f"cnn_research_survey_{time.strftime('%Y%m%d_%H%M%S')}.docx"
                st.success(tr("download_ready"))
            except Exception as exc:
                st.error(f"Không thể tạo báo cáo Word: {type(exc).__name__}: {exc}")
        if st.session_state.get("cnnlab_survey_docx"):
            st.download_button("Download Word report" if st.session_state.get("cnnlab_lang") == "en" else ("下載 Word 報告" if st.session_state.get("cnnlab_lang") == "zh-TW" else "Tải báo cáo Word"), data=st.session_state["cnnlab_survey_docx"], file_name=st.session_state.get("cnnlab_survey_name", "cnn_research_survey.docx"), mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        st.markdown(f"### {tr('results')}")
        df = pd.DataFrame(experiments)
        cols = [c for c in ["run_id", "model_name", "created_at", "epochs", "train_images_used", "val_accuracy", "val_f1", "val_roc_auc", "elapsed_seconds"] if c in df.columns]
        st.dataframe(df[cols], use_container_width=True, hide_index=True)
